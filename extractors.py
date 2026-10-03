"""업로드 파일 → 텍스트 추출 (PDF / PPTX / TXT / MD / 이미지) + OCR.

OCR 모드
- "off":    텍스트 레이어만 읽는다. 이미지는 건너뛴다.
- "auto":   텍스트 레이어가 없는 곳만 OCR (스캔 PDF 페이지, 이미지 파일, 슬라이드 속 그림).
- "always": PDF는 모든 페이지를 OCR (그림 속 글자·표까지 필요할 때).

OCR 엔진
- "openai":    OpenAI 비전 모델. 한글·의학용어·표 인식이 정확하지만 API 비용이 든다.
- "tesseract": 무료 로컬 OCR. tesseract-ocr(+kor) 설치 필요 (Streamlit Cloud는 packages.txt).
"""
from __future__ import annotations

import base64
import shutil
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

TEXT_TYPES = ["txt", "md"]
IMAGE_TYPES = ["png", "jpg", "jpeg", "webp", "bmp", "tif", "tiff"]
DOC_TYPES = ["pdf", "pptx"]
ALL_TYPES = DOC_TYPES + TEXT_TYPES + IMAGE_TYPES

OCR_MODES = ("off", "auto", "always")
OCR_ENGINES = ("openai", "tesseract")

_MIN_PAGE_CHARS = 30         # 이보다 글자가 적은 PDF 페이지는 스캔본으로 보고 OCR
_MIN_IMAGE_PX = 120          # 이보다 작은 슬라이드 그림(아이콘·로고)은 OCR 생략
_PDF_RENDER_SCALE = 2.0      # 72dpi × 2 ≈ 144dpi
_MAX_IMAGE_SIDE = 2000       # OpenAI로 보내기 전 긴 변 축소
_OCR_WORKERS = 4

_OCR_PROMPT = (
    "이 이미지에 있는 모든 글자를 빠짐없이 그대로 옮겨 적어라. "
    "한국어·영어·의학용어·약어·숫자·단위를 원문 그대로 유지하고, 번역하거나 요약하지 마라. "
    "표는 마크다운 표로, 목록은 줄바꿈을 유지해서 적어라. "
    "그림·도표 속 라벨도 포함하되 설명이나 해설은 붙이지 마라. "
    "글자가 없으면 빈 문자열만 출력하라."
)


@dataclass(frozen=True)
class OCRConfig:
    mode: str = "off"
    engine: str = "openai"
    model: str = "gpt-4o-mini"
    lang: str = "kor+eng"

    @property
    def enabled(self) -> bool:
        return self.mode != "off"


def tesseract_available() -> bool:
    try:
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return shutil.which("tesseract") is not None


# ─── 공개 API ───

def extract_text(file, ocr: OCRConfig | None = None) -> str:
    """Streamlit UploadedFile, 경로, 파일 객체 모두 받는다."""
    name = getattr(file, "name", str(file))
    if hasattr(file, "getvalue"):
        data = file.getvalue()          # read()와 달리 rerun 때도 빈 값이 되지 않음
    elif hasattr(file, "read"):
        data = file.read()
    else:
        data = Path(file).read_bytes()
    return extract_bytes(name, data, ocr)


def extract_bytes(name: str, data: bytes, ocr: OCRConfig | None = None) -> str:
    ocr = ocr or OCRConfig()
    ext = Path(name).suffix.lower().lstrip(".")
    if ext == "pdf":
        return _from_pdf(data, ocr)
    if ext == "pptx":
        return _from_pptx(data, ocr)
    if ext in TEXT_TYPES:
        return _decode_text(data)
    if ext in IMAGE_TYPES:
        if not ocr.enabled:
            return f"[이미지 {name}: OCR을 켜야 텍스트를 추출할 수 있습니다]"
        return _ocr_images([_load_image(data)], ocr)[0]
    return f"[지원 안 되는 형식: {name}]"


def estimate_tokens(text: str) -> int:
    """한글 글자당 약 0.7 토큰 (대략값)."""
    return int(len(text) * 0.7)


# ─── 형식별 추출 ───

def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp949"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _from_pdf(data: bytes, ocr: OCRConfig) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    texts = [(p.extract_text() or "").strip() for p in reader.pages]

    if ocr.mode == "always":
        targets = list(range(len(texts)))
    elif ocr.mode == "auto":
        targets = [i for i, t in enumerate(texts) if len("".join(t.split())) < _MIN_PAGE_CHARS]
    else:
        targets = []

    if targets:
        images = _render_pdf_pages(data, targets)
        for i, ocr_text in zip(targets, _ocr_images(images, ocr)):
            # 텍스트 레이어가 있던 페이지는 OCR 결과가 더 길 때만 교체 (그림 속 글자 포함)
            if len(ocr_text) > len(texts[i]):
                texts[i] = ocr_text

    return "\n\n".join(f"--- Page {i} ---\n{t}" for i, t in enumerate(texts, 1) if t)


def _render_pdf_pages(data: bytes, indices: list[int]) -> list:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    try:
        return [pdf[i].render(scale=_PDF_RENDER_SCALE).to_pil() for i in indices]
    finally:
        pdf.close()


def _from_pptx(data: bytes, ocr: OCRConfig) -> str:
    from pptx import Presentation

    prs = Presentation(BytesIO(data))
    slides_out = []
    pending = []   # (slide_idx, PIL.Image) — OCR은 모아서 병렬 처리
    for si, slide in enumerate(prs.slides):
        lines = []
        for shape in slide.shapes:
            _collect_shape(shape, lines, pending if ocr.enabled else None, si)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[발표자 노트] {notes}")
        slides_out.append(lines)

    if pending:
        for (si, _), text in zip(pending, _ocr_images([img for _, img in pending], ocr)):
            if text:
                slides_out[si].append(f"[그림 OCR] {text}")

    return "\n".join(
        "\n".join([f"--- Slide {i} ---", *lines]) for i, lines in enumerate(slides_out, 1)
    )


def _collect_shape(shape, lines: list, pending: list | None, slide_idx: int) -> None:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in shape.shapes:
            _collect_shape(sub, lines, pending, slide_idx)
        return
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        lines.append(shape.text_frame.text.strip())
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
    if pending is not None and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            img = _load_image(shape.image.blob)
        except Exception:
            return
        if min(img.size) >= _MIN_IMAGE_PX:
            pending.append((slide_idx, img))


# ─── OCR ───

def _load_image(data: bytes):
    from PIL import Image

    img = Image.open(BytesIO(data))
    img.load()
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    return img


def _ocr_images(images: list, ocr: OCRConfig) -> list[str]:
    if not images:
        return []
    if ocr.engine == "tesseract":
        fn = _ocr_tesseract
    elif ocr.engine == "openai":
        fn = _ocr_openai
    else:
        raise ValueError(f"알 수 없는 OCR 엔진: {ocr.engine}")

    def run(img):
        try:
            return fn(img, ocr).strip()
        except RuntimeError:
            raise                      # API 키 없음·Tesseract 미설치: 사용자에게 바로 알림
        except Exception as e:         # 페이지 하나 실패는 표시만 하고 계속
            return f"[OCR 실패: {e}]"

    if len(images) == 1:
        return [run(images[0])]
    with ThreadPoolExecutor(max_workers=_OCR_WORKERS) as pool:
        return list(pool.map(run, images))


def _ocr_tesseract(img, ocr: OCRConfig) -> str:
    if not tesseract_available():
        raise RuntimeError(
            "Tesseract가 설치되어 있지 않습니다. "
            "tesseract-ocr / tesseract-ocr-kor를 설치하거나 OCR 엔진을 OpenAI로 바꾸세요."
        )
    import pytesseract

    # psm 6(한 덩어리 문단): 기본값(psm 3)은 한글 줄을 조각내서 순서가 뒤섞인다
    return pytesseract.image_to_string(
        img.convert("L"), lang=ocr.lang, config="--psm 6 -c preserve_interword_spaces=1"
    )


def _ocr_openai(img, ocr: OCRConfig) -> str:
    from llm import completion_kwargs, get_client

    img = img.copy()
    img.thumbnail((_MAX_IMAGE_SIDE, _MAX_IMAGE_SIDE))
    buf = BytesIO()
    img.save(buf, format="PNG")
    url = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()

    resp = get_client().chat.completions.create(
        messages=[{
            "role": "user",
            "content": [
                {"type": "text", "text": _OCR_PROMPT},
                {"type": "image_url", "image_url": {"url": url, "detail": "high"}},
            ],
        }],
        **completion_kwargs(ocr.model, max_tokens=4096, temperature=0),
    )
    return resp.choices[0].message.content or ""
