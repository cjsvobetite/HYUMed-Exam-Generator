from io import BytesIO

import pytest
from PIL import Image, ImageDraw, ImageFont

import extractors
from extractors import OCRConfig, extract_bytes, extract_text, tesseract_available


def _text_image(text="HEART FAILURE", size=(900, 200)):
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 64)
    except OSError:
        font = ImageFont.load_default(size=64)
    draw.text((30, 50), text, fill="black", font=font)
    return img


def _png_bytes(img):
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _pdf(pages):
    """pages: str(텍스트 페이지) 또는 PIL.Image(스캔 페이지) 리스트."""
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buf = BytesIO()
    c = canvas.Canvas(buf)
    for page in pages:
        if isinstance(page, str):
            c.drawString(72, 720, page)
        else:
            c.drawImage(ImageReader(page), 72, 500, width=450, height=100)
        c.showPage()
    c.save()
    return buf.getvalue()


@pytest.fixture
def fake_ocr(monkeypatch):
    calls = []

    def fake(img, ocr):
        calls.append(img.size)
        return "OCR TEXT"
    monkeypatch.setattr(extractors, "_ocr_openai", fake)
    return calls


def test_text_utf8_and_cp949():
    assert extract_bytes("a.txt", "심근경색".encode("utf-8")) == "심근경색"
    assert extract_bytes("a.md", "심근경색".encode("cp949")) == "심근경색"


def test_uploaded_file_can_be_read_twice():
    f = BytesIO(b"hello")
    f.name = "a.txt"
    assert extract_text(f) == "hello"
    assert extract_text(f) == "hello"


def test_unsupported_type():
    assert "지원 안 되는 형식" in extract_bytes("a.docx", b"x")


def test_pdf_text_layer_without_ocr(fake_ocr):
    data = _pdf(["Myocardial infarction is necrosis of heart muscle"])
    out = extract_bytes("a.pdf", data, OCRConfig(mode="auto"))
    assert "Myocardial infarction" in out
    assert fake_ocr == []          # 텍스트가 있으니 OCR 안 함


def test_pdf_auto_ocr_only_scanned_pages(fake_ocr):
    data = _pdf(["Myocardial infarction is necrosis of heart muscle", _text_image()])
    out = extract_bytes("a.pdf", data, OCRConfig(mode="auto"))
    assert "Myocardial infarction" in out
    assert "--- Page 2 ---\nOCR TEXT" in out
    assert len(fake_ocr) == 1


def test_pdf_ocr_off_skips_scanned_pages(fake_ocr):
    data = _pdf([_text_image()])
    assert extract_bytes("a.pdf", data, OCRConfig(mode="off")) == ""
    assert fake_ocr == []


def test_pdf_always_ocrs_every_page(fake_ocr):
    data = _pdf(["MI", "CHF"])
    out = extract_bytes("a.pdf", data, OCRConfig(mode="always"))
    assert len(fake_ocr) == 2
    assert out.count("OCR TEXT") == 2


def test_image_requires_ocr(fake_ocr):
    data = _png_bytes(_text_image())
    assert "OCR을 켜야" in extract_bytes("a.png", data)
    assert extract_bytes("a.png", data, OCRConfig(mode="auto")) == "OCR TEXT"


def test_ocr_page_failure_is_reported_inline(monkeypatch):
    def boom(img, ocr):
        raise ValueError("bad image")
    monkeypatch.setattr(extractors, "_ocr_openai", boom)
    out = extract_bytes("a.png", _png_bytes(_text_image()), OCRConfig(mode="auto"))
    assert out == "[OCR 실패: bad image]"


def test_missing_api_key_is_raised(monkeypatch):
    import llm
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    llm.get_client.cache_clear()
    with pytest.raises(llm.MissingAPIKeyError):
        extract_bytes("a.png", _png_bytes(_text_image()), OCRConfig(mode="auto"))
    llm.get_client.cache_clear()


def test_pptx_text_table_group_notes_and_pictures(fake_ocr):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Heart failure"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Drug", "Class"
    table.cell(1, 0).text, table.cell(1, 1).text = "Furosemide", "Loop"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(4), Inches(2), Inches(1)).text_frame.text = "Grouped text"
    slide.shapes.add_picture(BytesIO(_png_bytes(_text_image())), Inches(1), Inches(5))
    slide.shapes.add_picture(BytesIO(_png_bytes(Image.new("RGB", (40, 40)))), Inches(6), Inches(5))  # 아이콘
    slide.notes_slide.notes_text_frame.text = "Speaker note"
    buf = BytesIO()
    prs.save(buf)

    no_ocr = extract_bytes("a.pptx", buf.getvalue())
    assert "--- Slide 1 ---" in no_ocr
    assert "Heart failure" in no_ocr
    assert "| Furosemide | Loop |" in no_ocr
    assert "Grouped text" in no_ocr
    assert "[발표자 노트] Speaker note" in no_ocr
    assert "OCR" not in no_ocr and fake_ocr == []

    with_ocr = extract_bytes("a.pptx", buf.getvalue(), OCRConfig(mode="auto"))
    assert "[그림 OCR] OCR TEXT" in with_ocr
    assert len(fake_ocr) == 1      # 작은 아이콘은 건너뜀


@pytest.mark.skipif(not tesseract_available(), reason="tesseract 미설치")
def test_tesseract_reads_scanned_pdf():
    data = _pdf([_text_image("HEART FAILURE")])
    out = extract_bytes("a.pdf", data, OCRConfig(mode="auto", engine="tesseract", lang="eng"))
    assert "HEART" in out.upper()


@pytest.mark.skipif(not tesseract_available(), reason="tesseract 미설치")
def test_tesseract_reads_korean():
    import glob
    fonts = glob.glob("/usr/share/fonts/**/NanumGothic.ttf", recursive=True)
    if not fonts:
        pytest.skip("한글 폰트 없음")
    img = Image.new("RGB", (1200, 300), "white")
    ImageDraw.Draw(img).text((40, 100), "심부전 치료에는 이뇨제를 사용한다", fill="black",
                             font=ImageFont.truetype(fonts[0], 48))
    out = extract_bytes("a.png", _png_bytes(img), OCRConfig(mode="auto", engine="tesseract"))
    assert "심부전 치료에는 이뇨제를" in out
